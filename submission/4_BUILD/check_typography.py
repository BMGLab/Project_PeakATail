"""Read the built .docx back and assert the PI typography spec on RESOLVED values.

A run with no explicit size inherits from its paragraph style, then that style's
ancestors, then docDefaults -- so checking run.font.size alone proves nothing.
"""
import sys, re, collections
import docx
from docx.shared import Pt

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
SPEC = dict(body=11.0, heading=12.0, legend=10.0, table=9.0)
LEGEND_RE = re.compile(r"^\s*(?:Fig(?:ure)?|Table)\s+S?\d+\s*[|.:]")


def doc_default_pt(d):
    dd = d.styles.element.find(W + "docDefaults")
    sz = dd.find(W + "rPrDefault").find(W + "rPr").find(W + "sz")
    return int(sz.get(W + "val")) / 2.0


def style_chain(style):
    seen = []
    while style is not None and style.name not in seen:
        yield style
        seen.append(style.name)
        style = getattr(style, "base_style", None)


def resolved_pt(run, par, default):
    if run.font.size is not None:
        return run.font.size.pt
    for st in style_chain(par.style):
        if st.font.size is not None:
            return st.font.size.pt
    return default


def resolved_bold(par):
    for st in style_chain(par.style):
        if st.font.bold is not None:
            return st.font.bold
    return False


def resolved_jc(par):
    for st in style_chain(par.style):
        p = st.element.find(W + "pPr")
        if p is not None and p.find(W + "jc") is not None:
            return p.find(W + "jc").get(W + "val")
    return None


def audit(path):
    d = docx.Document(path)
    default = doc_default_pt(d)
    tbl_par_ids = {id(p._p) for t in d.tables for r in t.rows for c in r.cells
                   for p in c.paragraphs}
    buckets = collections.defaultdict(collections.Counter)
    unjust = 0
    body_n = 0
    for par in d.paragraphs:
        if not par.text.strip():
            continue
        sizes = {resolved_pt(r, par, default) for r in par.runs} or {default}
        if par.style.name.startswith("Heading"):
            kind = "heading"
        elif par.style.name == "FigureLegend" or LEGEND_RE.match(par.text.lstrip("*")):
            kind = "legend"
        elif par.style.name in ("Body Text", "First Paragraph", "Normal", "Compact"):
            kind = "body"
            body_n += 1
            if resolved_jc(par) != "both":
                unjust += 1
        else:
            kind = "other:" + par.style.name
        for s in sizes:
            buckets[kind][s] += 1
    for t in d.tables:
        for row in t.rows:
            for cell in row.cells:
                for par in cell.paragraphs:
                    if par.text.strip():
                        for r in par.runs:
                            buckets["table"][resolved_pt(r, par, default)] += 1

    print(f"\n=== {path.split('/')[-1]}  (docDefaults {default} pt) ===")
    fails = []
    for kind in ("body", "heading", "legend", "table"):
        c = buckets.get(kind)
        if not c:
            print(f"  {kind:9s} : (none present)")
            continue
        want = SPEC[kind]
        off = {k: v for k, v in c.items() if abs(k - want) > 0.01}
        flag = "OK  " if not off else "FAIL"
        if off:
            fails.append(f"{kind}: {dict(sorted(off.items()))} (want {want} pt)")
        print(f"  {flag} {kind:9s}: {dict(sorted(c.items()))}  want {want} pt")
    hb = [p.style.name for p in d.paragraphs
          if p.style.name.startswith("Heading") and p.text.strip() and not resolved_bold(p)]
    print(f"  {'OK  ' if not hb else 'FAIL'} headings bold : "
          f"{len(hb)} not bold" + (f" -> {sorted(set(hb))}" if hb else ""))
    if hb:
        fails.append(f"headings not bold: {sorted(set(hb))}")
    print(f"  {'OK  ' if not unjust else 'FAIL'} justified     : "
          f"{body_n - unjust}/{body_n} body paragraphs justified")
    if unjust:
        fails.append(f"{unjust} body paragraphs not justified")
    extra = {k: dict(v) for k, v in buckets.items() if k.startswith("other:")}
    if extra:
        print(f"  note  other styles present: {extra}")
    return fails


bad = []
for f in sys.argv[1:]:
    bad += [f"{f.split('/')[-1]}: {x}" for x in audit(f)]
print()
if bad:
    print("SPEC VIOLATIONS:")
    for b in bad:
        print("  -", b)
    sys.exit(1)
print("all three documents match the PI typography spec")
