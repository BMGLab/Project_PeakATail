#!/usr/bin/env python3
"""Finish a pandoc-produced .docx: continuous line numbers, page-number footer,
document properties.  Word features pandoc 2.9's docx writer cannot emit.

    python3 postprocess_docx.py <file.docx> [--no-line-numbers] [--title "..."]

Content is never touched — only section properties, the footer, and docProps.
"""
import sys

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

# CT_SectPr child order (ECMA-376); we insert lnNumType / pgNumType in place.
SECTPR_ORDER = [
    "headerReference", "footerReference", "footnotePr", "endnotePr", "type",
    "pgSz", "pgMar", "paperSrc", "pgBorders", "lnNumType", "pgNumType", "cols",
    "formProt", "vAlign", "noEndnote", "titlePg", "textDirection", "bidi",
    "rtlGutter", "docGrid", "printerSettings", "sectPrChange",
]


def insert_ordered(sectPr, el):
    name = el.tag.split("}")[1]
    pos = SECTPR_ORDER.index(name)
    for child in sectPr:
        cname = child.tag.split("}")[1]
        if cname in SECTPR_ORDER and SECTPR_ORDER.index(cname) > pos:
            child.addprevious(el)
            return
    sectPr.append(el)


def add_line_numbers(section, distance=360):
    sectPr = section._sectPr
    for old in sectPr.findall(qn("w:lnNumType")):
        sectPr.remove(old)
    ln = OxmlElement("w:lnNumType")
    ln.set(qn("w:countBy"), "1")
    # NB: no w:start.  Word defaults it to 1; LibreOffice ADDS w:start to the
    # count, so w:start="1" makes LibreOffice label the first line "2".
    ln.set(qn("w:restart"), "continuous")   # continuous through the document
    ln.set(qn("w:distance"), str(distance))
    insert_ordered(sectPr, ln)


def add_page_number_start(section):
    sectPr = section._sectPr
    for old in sectPr.findall(qn("w:pgNumType")):
        sectPr.remove(old)
    pn = OxmlElement("w:pgNumType")
    pn.set(qn("w:start"), "1")
    insert_ordered(sectPr, pn)


def field_run(p, instr):
    """Append a Word field (e.g. ' PAGE ') to paragraph p."""
    def run(children):
        r = OxmlElement("w:r")
        for c in children:
            r.append(c)
        p._p.append(r)
        return r

    b = OxmlElement("w:fldChar")
    b.set(qn("w:fldCharType"), "begin")
    run([b])
    it = OxmlElement("w:instrText")
    it.set(qn("xml:space"), "preserve")
    it.text = instr
    run([it])
    s = OxmlElement("w:fldChar")
    s.set(qn("w:fldCharType"), "separate")
    run([s])
    t = OxmlElement("w:t")
    t.text = "1"
    run([t])
    e = OxmlElement("w:fldChar")
    e.set(qn("w:fldCharType"), "end")
    run([e])


def add_footer_page_numbers(section, font="Times New Roman", pt=10):
    footer = section.footer
    footer.is_linked_to_previous = False
    p = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
    for r in list(p.runs):
        r._r.getparent().remove(r._r)
    p.text = ""
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    # keep the footer out of the line-number sequence (LibreOffice otherwise
    # counts it, pushing the body's first line to number 2)
    pPr = p._p.get_or_add_pPr()
    if pPr.find(qn("w:suppressLineNumbers")) is None:
        pPr.insert(0, OxmlElement("w:suppressLineNumbers"))
    r0 = p.add_run("Page ")
    field_run(p, " PAGE ")
    r1 = p.add_run(" of ")
    field_run(p, " NUMPAGES ")
    from docx.shared import Pt
    for r in p.runs:
        r.font.name = font
        r.font.size = Pt(pt)
    # field runs are raw XML; give them the same rPr
    for r in p._p.findall(qn("w:r")):
        if r.find(qn("w:rPr")) is None:
            rPr = OxmlElement("w:rPr")
            rf = OxmlElement("w:rFonts")
            for a in ("w:ascii", "w:hAnsi", "w:cs"):
                rf.set(qn(a), font)
            sz = OxmlElement("w:sz")
            sz.set(qn("w:val"), str(int(pt * 2)))
            rPr.append(rf)
            rPr.append(sz)
            r.insert(0, rPr)
    return p


def main():
    path = sys.argv[1]
    args = sys.argv[2:]
    line_numbers = "--no-line-numbers" not in args
    title = None
    if "--title" in args:
        title = args[args.index("--title") + 1]

    d = docx.Document(path)
    for sec in d.sections:
        if line_numbers:
            add_line_numbers(sec)
        add_page_number_start(sec)
        add_footer_page_numbers(sec)

    cp = d.core_properties
    if title:
        cp.title = title
    cp.author = "Amir Amiri Tabat; Yasin Kaymaz"
    cp.comments = ("Built from manuscript/00_draft by submission/build/build.sh; "
                   "content unaltered.")
    d.save(path)
    print("postprocessed %s  (line numbers: %s; footer: 'Page X of Y'; "
          "%d section(s))" % (path, "continuous" if line_numbers else "off",
                              len(d.sections)))


if __name__ == "__main__":
    main()
